# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document


SOURCE = Path("work-docx/0.docx")
OUTPUT = Path("work-docx/0_filled_bitcoin_indexer.docx")


def replace_paragraph(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


doc = Document(SOURCE)

replace_paragraph(
    doc.paragraphs[31],
    "Выполнить анализ предметной области индексирования данных Bitcoin и "
    "ограничений прямого доступа через Bitcoin Core RPC. Спроектировать и "
    "реализовать систему Bitcoin Blockchain Indexer для получения данных от "
    "узла Bitcoin Core, хранения блоков, транзакций, UTXO и балансов в "
    "PostgreSQL, обработки mempool и reorg, управления заданиями индексирования "
    "и выдачи данных через REST API.",
)

stage_rows = [
    [
        "1",
        "Анализ предметной области и аналогов",
        "15%",
        "20.02.2025-\n05.03.2025",
        "",
    ],
    [
        "2",
        "Требования, архитектура и модель данных",
        "25%",
        "05.03.2025-\n25.03.2025",
        "",
    ],
    [
        "3",
        "Реализация RPC-клиента, индексатора, API и jobs",
        "40%",
        "25.03.2025-\n30.04.2025",
        "",
    ],
    [
        "4",
        "Тестирование и оформление ВКР",
        "20%",
        "30.04.2025-\n18.05.2025",
        "",
    ],
]

for row, values in zip(doc.tables[0].rows[1:], stage_rows):
    for cell, value in zip(row.cells, values):
        replace_paragraph(cell.paragraphs[0], value)
        for paragraph in cell.paragraphs[1:]:
            replace_paragraph(paragraph, "")

sources = [
    "5. Исходные материалы и пособия\n"
    "1. Nakamoto S. Bitcoin: A Peer-to-Peer Electronic Cash System. — 2008.",
    "2. Narayanan A., Bonneau J., Felten E., Miller A., Goldfeder S. "
    "Bitcoin and Cryptocurrency Technologies. — Princeton University Press, 2016.",
    "3. Antonopoulos A.M. Mastering Bitcoin: Programming the Open Blockchain. "
    "— 2nd ed. — O'Reilly Media, 2017.",
    "4. Kleppmann M. Designing Data-Intensive Applications. — O'Reilly Media, 2017.",
    "5. Bitcoin Core RPC Documentation. — URL: https://developer.bitcoin.org/reference/rpc/; "
    "PostgreSQL Documentation. — URL: https://www.postgresql.org/docs/.",
]

for index, text in zip(range(35, 40), sources):
    replace_paragraph(doc.paragraphs[index], text)

doc.save(OUTPUT)
print(OUTPUT)
